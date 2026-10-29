"""Integración: la API real contra el emulador de Firestore y Auth (sin tocar producción).

    firebase emulators:start --only firestore,auth --project fusa-conecta
    FIRESTORE_EMULATOR_HOST=127.0.0.1:8099 FIREBASE_AUTH_EMULATOR_HOST=127.0.0.1:9099 pytest -m integracion

Se salta sola si el emulador no está configurado.
"""

from __future__ import annotations

import os
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from fastapi.testclient import TestClient

pytestmark = [
    pytest.mark.integracion,
    pytest.mark.skipif(not os.environ.get("FIRESTORE_EMULATOR_HOST"), reason="sin emulador de Firestore"),
]
PROYECTO = os.environ.get("FIREBASE_PROJECT", "fusa-conecta")  # en CI: demo-fusa-conecta (solo emulador)


def _fs(ruta: str) -> str:
    return f"http://{os.environ['FIRESTORE_EMULATOR_HOST']}/v1/projects/{PROYECTO}/databases/(default)/documents/{ruta}"


def _valor(v):
    if v is None:
        return {"nullValue": None}
    if isinstance(v, bool):
        return {"booleanValue": v}
    if isinstance(v, int):
        return {"integerValue": str(v)}
    if isinstance(v, float):
        return {"doubleValue": v}
    if isinstance(v, datetime):
        return {"timestampValue": v.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}
    if isinstance(v, list):
        return {"arrayValue": {"values": [_valor(x) for x in v]}}
    return {"stringValue": str(v)}


def poner(ruta: str, datos: dict) -> None:
    r = httpx.patch(
        _fs(ruta), json={"fields": {k: _valor(v) for k, v in datos.items()}}, headers={"Authorization": "Bearer owner"}
    )
    r.raise_for_status()


def cuenta(correo: str) -> tuple[str, str]:
    """Crea (o entra con) una cuenta en el emulador de Auth y devuelve (uid, idToken)."""
    base = f"http://{os.environ.get('FIREBASE_AUTH_EMULATOR_HOST', '127.0.0.1:9099')}/identitytoolkit.googleapis.com/v1/accounts"
    datos = {"email": correo, "password": "PruebaApi2026", "returnSecureToken": True}
    r = httpx.post(base + ":signUp?key=demo", json=datos)
    if r.status_code != 200:
        r = httpx.post(base + ":signInWithPassword?key=demo", json=datos)
    j = r.json()
    return j["localId"], j["idToken"]


@pytest.fixture(scope="module")
def entorno():
    os.environ.setdefault("FIREBASE_AUTH_EMULATOR_HOST", "127.0.0.1:9099")
    os.environ["CACHE_SEGUNDOS"] = "0"
    from app import main
    from app.repositorio import usar_repositorio

    usar_repositorio(None)
    duena, token_duena = cuenta("duena-api@fusa.test")
    admin, token_admin = cuenta("admin-api@fusa.test")
    _, token_otro = cuenta("otro-api@fusa.test")
    ahora, mes = datetime.now(timezone.utc), datetime.now(timezone(timedelta(hours=-5))).strftime("%Y-%m")
    poner(f"usuarios/{admin}", {"rol": "admin", "nombre": "Admin API"})
    poner(f"usuarios/{duena}", {"rol": "usuario", "nombre": "Dueña API"})
    poner(
        "comercios/api-panaderia",
        {
            "nombre": "Panadería de Prueba API",
            "cat": "comercio",
            "sub": "panaderias",
            "estado_publicacion": "aprobado",
            "propietario_uid": duena,
            "horario": "Lun-Dom 6:00-20:00",
            "municipio": "Fusagasugá",
            "esPremium": True,
        },
    )
    poner(
        "comercios/api-pendiente",
        {
            "nombre": "Negocio en revisión API",
            "cat": "comercio",
            "estado_publicacion": "pendiente",
            "propietario_uid": duena,
        },
    )
    poner(f"comercios/api-panaderia/estadisticas/{mes}", {"vistas": 42, "toques_wa": 7, "toques_llamar": 2})
    poner(
        "canjes/api-k1",
        {
            "cupon_id": "api-c1",
            "comercio_id": "api-panaderia",
            "estado": "usado",
            "fecha_emision": ahora,
            "fecha_canje": ahora,
        },
    )
    poner(
        "eventos/api-e1",
        {
            "titulo": "Evento de integración API",
            "categoria": "musica",
            "estado": "publicado",
            "lugar_nombre": "Concha acústica",
            "fecha_inicio": ahora + timedelta(days=1),
            "fecha_fin": ahora + timedelta(days=1, hours=2),
            "tipo": "plan",
        },
    )
    poner(
        "eventos/api-e2",
        {
            "titulo": "Evento pendiente API",
            "categoria": "musica",
            "estado": "pendiente",
            "lugar_nombre": "x",
            "fecha_inicio": ahora + timedelta(days=1),
            "fecha_fin": ahora + timedelta(days=1, hours=2),
        },
    )
    main.limiter.reset()
    yield TestClient(main.app), token_duena, token_admin, token_otro, mes
    usar_repositorio(None)


def test_salud_con_firestore(entorno):
    c, *_ = entorno
    assert c.get("/v1/salud").json()["firestore"] is True


def test_solo_lo_publico(entorno):
    c, *_ = entorno
    ids = [x["id"] for x in c.get("/v1/comercios", params={"q": "API", "por_pagina": 50}).json()["resultados"]]
    assert "api-panaderia" in ids and "api-pendiente" not in ids
    eventos = [e["id"] for e in c.get("/v1/eventos").json()["resultados"]]
    assert "api-e1" in eventos and "api-e2" not in eventos


def test_reporte_con_token_real_del_emulador(entorno):
    c, duena, admin, otro, mes = entorno
    assert (
        c.get("/v1/comercios/api-panaderia/reporte", headers={"Authorization": "Bearer token-falso"}).status_code == 401
    )
    assert c.get("/v1/comercios/api-panaderia/reporte", headers={"Authorization": f"Bearer {otro}"}).status_code == 403
    r = c.get(
        "/v1/comercios/api-panaderia/reporte",
        params={"formato": "xlsx", "mes": mes},
        headers={"Authorization": f"Bearer {duena}"},
    )
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from io import BytesIO

    from openpyxl import load_workbook

    filas = {f[0]: f[1] for f in load_workbook(BytesIO(r.content))["Resumen"].iter_rows(min_row=5, values_only=True)}
    assert filas["Vistas de la ficha"] == 42 and filas["Toques en WhatsApp"] == 7 and filas["Cupones canjeados"] == 1
    assert c.get(
        "/v1/comercios/api-panaderia/reporte", params={"formato": "pdf"}, headers={"Authorization": f"Bearer {admin}"}
    ).content.startswith(b"%PDF")


def test_estadisticas_admin(entorno):
    c, duena, admin, *_ = entorno
    assert c.get("/v1/admin/estadisticas", headers={"Authorization": f"Bearer {duena}"}).status_code == 403
    r = c.get("/v1/admin/estadisticas", headers={"Authorization": f"Bearer {admin}"}).json()
    assert (
        r["negocios_aprobados"] >= 1
        and r["negocios_pendientes"] >= 1
        and r["usuarios"] >= 2
        and r["cupones_canjeados"] >= 1
    )


def test_aviso_de_aprobacion_contra_el_emulador(entorno):
    """La novedad queda en usuarios/{uid}/notificaciones y el token inválido se borra (FCM simulado)."""
    from app import avisos, main

    c, duena_token, admin, *_ = entorno
    uid = _uid(duena_token)
    poner(f"usuarios/{uid}/dispositivos/tok-bueno", {"plataforma": "android"})
    poner(f"usuarios/{uid}/dispositivos/tok-viejo", {"plataforma": "android"})

    class Simulado:
        enviados: list = []

        def enviar(self, tokens, titulo, texto, enlace, etiqueta):
            self.enviados.append((sorted(tokens), texto))
            return 1, ["tok-viejo"]

    fcm = Simulado()
    avisos._estado.clear()
    main.app.dependency_overrides[main._avisador] = lambda: fcm
    try:
        cuerpo = {"tipo": "comercio", "id": "api-panaderia", "resultado": "aprobado"}
        assert (
            c.post(
                "/v1/notificaciones/aprobacion", json=cuerpo, headers={"Authorization": f"Bearer {duena_token}"}
            ).status_code
            == 403
        )
        r = c.post("/v1/notificaciones/aprobacion", json=cuerpo, headers={"Authorization": f"Bearer {admin}"})
        assert r.status_code == 200 and r.json()["enviados"] == 1
        assert fcm.enviados == [
            (["tok-bueno", "tok-viejo"], "¡Tu negocio Panadería de Prueba API ya está publicado en Fusa Conecta!")
        ]
        nov = httpx.get(
            _fs(f"usuarios/{uid}/notificaciones/comercio-api-panaderia-aprobado"),
            headers={"Authorization": "Bearer owner"},
        )
        assert nov.status_code == 200 and nov.json()["fields"]["enlace"]["stringValue"] == "#/mi-negocio"
        assert (
            httpx.get(
                _fs(f"usuarios/{uid}/dispositivos/tok-viejo"), headers={"Authorization": "Bearer owner"}
            ).status_code
            == 404
        )
        assert (
            httpx.get(
                _fs(f"usuarios/{uid}/dispositivos/tok-bueno"), headers={"Authorization": "Bearer owner"}
            ).status_code
            == 200
        )
        # El reintento no repite el aviso
        assert (
            c.post("/v1/notificaciones/aprobacion", json=cuerpo, headers={"Authorization": f"Bearer {admin}"}).json()[
                "repetido"
            ]
            is True
        )
        assert len(fcm.enviados) == 1
    finally:
        main.app.dependency_overrides.pop(main._avisador, None)


def _uid(token: str) -> str:
    import base64
    import json

    carga = token.split(".")[1]
    return json.loads(base64.urlsafe_b64decode(carga + "=" * (-len(carga) % 4)))["user_id"]
