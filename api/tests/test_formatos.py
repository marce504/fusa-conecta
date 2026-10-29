"""Exportaciones: GPX válido, ICS que cumple RFC 5545 y reportes Excel/PDF que se pueden abrir."""

import io
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from app import formatos

GPX_NS = "{http://www.topografix.com/GPX/1/1}"


def test_gpx_valido_con_alturas():
    texto = formatos.gpx(
        "Sendero <Cara del Indio>", [(4.1, -74.5, 1700.0), (4.2, -74.6, None)], "3,98 km", "OpenStreetMap"
    )
    raiz = ET.fromstring(texto.encode("utf-8"))
    puntos = raiz.findall(f".//{GPX_NS}trkpt")
    assert raiz.get("version") == "1.1" and len(puntos) == 2
    assert puntos[0].get("lat") == "4.100000" and puntos[0].find(f"{GPX_NS}ele").text == "1700.0"
    assert puntos[1].find(f"{GPX_NS}ele") is None
    assert raiz.find(f"{GPX_NS}trk/{GPX_NS}name").text == "Sendero <Cara del Indio>"  # escapado y legible


def test_ics_rfc5545():
    ev = {
        "id": "e1",
        "titulo": "Día Dulce, con dulces; y más",
        "ini": 1793473200000,
        "fin": 1793484000000,
        "lugar_nombre": "Estadio",
        "municipio": "Fusagasugá",
        "descripcion": "x" * 200,
        "categoria": "cultura",
        "aviso": False,
        "enlace_fuente": "https://a.gov.co/n",
    }
    texto = formatos.ics([ev, {**ev, "id": "e2", "aviso": True}])
    lineas = texto.split("\r\n")
    assert lineas[0] == "BEGIN:VCALENDAR" and texto.endswith("END:VCALENDAR\r\n")
    assert all(len(x.encode("utf-8")) <= 75 for x in lineas), "líneas plegadas a 75 octetos"
    assert "SUMMARY:Día Dulce\\, con dulces\\; y más" in texto and "SUMMARY:Aviso: Día Dulce" in texto
    assert texto.count("BEGIN:VEVENT") == 2 and "UID:e1@fusa-conecta.web.app" in texto
    desplegado = texto.replace("\r\n ", "")
    assert "x" * 200 in desplegado


REPORTE = {
    "negocio": "Panadería La Espiga",
    "mes": "2026-10",
    "generado": "2026-10-28 18:00",
    "cifras": {"vistas": 120, "toques_wa": 14, "toques_llamar": 3, "cupones_obtenidos": 2, "cupones_canjeados": 1},
    "cupones": [
        {
            "id": "c",
            "descripcion": "2x1",
            "estado": "activo",
            "cupo_total": 50,
            "cupo_restante": 40,
            "vence": "2026-11-07",
            "obtenidos": 2,
            "canjeados": 1,
        }
    ],
}


def test_reporte_excel():
    libro = load_workbook(io.BytesIO(formatos.reporte_xlsx(REPORTE)))
    assert libro.sheetnames == ["Resumen", "Cupones"]
    resumen = {f[0]: f[1] for f in libro["Resumen"].iter_rows(min_row=5, values_only=True)}
    assert (
        resumen["Vistas de la ficha"] == 120
        and resumen["Toques en WhatsApp"] == 14
        and resumen["Cupones canjeados"] == 1
    )
    assert list(libro["Cupones"].iter_rows(min_row=2, values_only=True))[0][:2] == ("2x1", "activo")


def test_reporte_pdf():
    pdf = formatos.reporte_pdf(REPORTE)
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500
    assert formatos.reporte_pdf({**REPORTE, "cupones": []}).startswith(b"%PDF")
