"""Exportaciones: GPX (rutas), ICS (agenda) y reportes del negocio en Excel (openpyxl) y PDF (ReportLab)."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Iterable
from xml.sax.saxutils import escape

SITIO = "https://fusa-conecta.web.app"
VIOLETA = "6B4EFF"


# ---------- GPX 1.1 ----------
def gpx(
    nombre: str, puntos: Iterable[tuple[float, float, float | None]], descripcion: str = "", fuente: str = ""
) -> str:
    """Trazado [(lat, lng, alt)] → GPX 1.1 que abren Wikiloc, Strava, Komoot, OsmAnd y Google Earth."""
    filas = []
    for lat, lng, alt in puntos:
        ele = f"<ele>{round(alt, 1)}</ele>" if alt is not None else ""
        filas.append(f'      <trkpt lat="{lat:.6f}" lon="{lng:.6f}">{ele}</trkpt>')
    meta = f"<name>{escape(nombre)}</name>" + (f"<desc>{escape(descripcion)}</desc>" if descripcion else "")
    autor = f"<author><name>{escape(fuente)}</name></author>" if fuente else ""
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<gpx version="1.1" creator="Fusa Conecta (https://fusa-conecta.web.app)" '
        'xmlns="http://www.topografix.com/GPX/1/1">\n'
        f'  <metadata>{meta}{autor}<link href="{SITIO}"><text>Fusa Conecta</text></link></metadata>\n'
        f"  <trk>\n    <name>{escape(nombre)}</name>\n    <trkseg>\n"
        + "\n".join(filas)
        + "\n    </trkseg>\n  </trk>\n</gpx>\n"
    )


# ---------- iCalendar (RFC 5545) ----------
def _ics_texto(s: str) -> str:
    return (
        str(s or "")
        .replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\r\n", "\\n")
        .replace("\n", "\\n")
    )


def _ics_fecha(ms: int) -> str:
    return datetime.fromtimestamp(ms / 1000, timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _plegar(linea: str) -> str:
    """Líneas de máximo 75 octetos, como pide el estándar (los calendarios cortan las largas)."""
    datos, salida = linea.encode("utf-8"), []
    while len(datos) > 75:
        corte = 75 if not salida else 74
        while corte > 0 and (datos[corte] & 0xC0) == 0x80:  # no partir una letra con tilde
            corte -= 1
        salida.append(datos[:corte].decode("utf-8"))
        datos = datos[corte:]
    salida.append(datos.decode("utf-8"))
    return "\r\n ".join(salida)


def ics(eventos: list[dict], nombre: str = "Fusa Conecta · Qué hacer en Fusa") -> str:
    """Calendario suscribible: Google Calendar, Apple Calendario y Outlook lo actualizan solos."""
    ahora = _ics_fecha(int(datetime.now(timezone.utc).timestamp() * 1000))
    lineas = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Fusa Conecta//Agenda//ES",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:" + _ics_texto(nombre),
        "X-WR-TIMEZONE:America/Bogota",
        "X-WR-CALDESC:Eventos y avisos de Fusagasugá y Pasca publicados en Fusa Conecta",
        "REFRESH-INTERVAL;VALUE=DURATION:PT6H",
        "X-PUBLISHED-TTL:PT6H",
    ]
    for e in eventos:
        titulo = ("Aviso: " if e.get("aviso") else "") + e["titulo"]
        desc = (e.get("descripcion") or "") + ("\n" if e.get("descripcion") else "") + f"{SITIO}/#/evento/{e['id']}"
        lineas += [
            "BEGIN:VEVENT",
            f"UID:{e['id']}@fusa-conecta.web.app",
            "DTSTAMP:" + ahora,
            "DTSTART:" + _ics_fecha(e["ini"]),
            "DTEND:" + _ics_fecha(e["fin"]),
            "SUMMARY:" + _ics_texto(titulo),
            "LOCATION:" + _ics_texto(f"{e.get('lugar_nombre', '')}, {e.get('municipio') or 'Fusagasugá'}"),
            "DESCRIPTION:" + _ics_texto(desc),
            "CATEGORIES:" + _ics_texto(e.get("categoria") or "otro"),
        ]
        if e.get("enlace_fuente", "").startswith("https://"):
            lineas.append("URL:" + e["enlace_fuente"])
        lineas.append("END:VEVENT")
    lineas.append("END:VCALENDAR")
    return "\r\n".join(_plegar(x) for x in lineas) + "\r\n"


# ---------- Reporte mensual del negocio ----------
CIFRAS = [
    ("vistas", "Vistas de la ficha"),
    ("toques_wa", "Toques en WhatsApp"),
    ("toques_llamar", "Toques en Llamar"),
    ("cupones_obtenidos", "Cupones obtenidos"),
    ("cupones_canjeados", "Cupones canjeados"),
]


def reporte_xlsx(r: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    libro = Workbook()
    h = libro.active
    h.title = "Resumen"
    titulo = Font(bold=True, size=14, color=VIOLETA)
    negrita, cabecera = Font(bold=True), PatternFill("solid", fgColor="EEEAFF")
    h["A1"], h["A1"].font = f"Reporte de {r['negocio']}", titulo
    h["A2"] = f"Mes: {r['mes']} · generado el {r['generado']} · Fusa Conecta"
    h.append([])
    h.append(["Cifra", "Total del mes"])
    for c in h[4]:
        c.font, c.fill = negrita, cabecera
    for clave, texto in CIFRAS:
        h.append([texto, r["cifras"].get(clave, 0)])
    h.column_dimensions["A"].width, h.column_dimensions["B"].width = 30, 16
    c = libro.create_sheet("Cupones")
    c.append(["Cupón", "Estado", "Cupo total", "Quedan", "Obtenidos en el mes", "Canjeados en el mes", "Vence"])
    for celda in c[1]:
        celda.font, celda.fill = negrita, cabecera
        celda.alignment = Alignment(wrap_text=True)
    for x in r["cupones"]:
        c.append(
            [
                x["descripcion"],
                x["estado"],
                x["cupo_total"],
                x["cupo_restante"],
                x["obtenidos"],
                x["canjeados"],
                x["vence"],
            ]
        )
    for col, ancho in zip("ABCDEFG", (34, 12, 11, 10, 14, 14, 12)):
        c.column_dimensions[col].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def reporte_pdf(r: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    salida = io.BytesIO()
    doc = SimpleDocTemplate(
        salida,
        pagesize=letter,
        title=f"Reporte {r['negocio']} {r['mes']}",
        author="Fusa Conecta",
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    est = getSampleStyleSheet()
    est["Title"].textColor = colors.HexColor("#" + VIOLETA)
    estilo_tabla = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEAFF")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9E3")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]
    )
    partes = [
        Paragraph(f"Reporte de {escape(r['negocio'])}", est["Title"]),
        Paragraph(f"Mes {r['mes']} · generado el {r['generado']} · Fusa Conecta", est["Normal"]),
        Spacer(1, 0.6 * cm),
        Table(
            [["Cifra", "Total del mes"]] + [[t, r["cifras"].get(k, 0)] for k, t in CIFRAS],
            colWidths=[10 * cm, 5 * cm],
            style=estilo_tabla,
        ),
        Spacer(1, 0.8 * cm),
        Paragraph("Cupones", est["Heading2"]),
    ]
    if r["cupones"]:
        filas = [["Cupón", "Estado", "Quedan", "Obtenidos", "Canjeados"]] + [
            [
                Paragraph(escape(x["descripcion"]), est["Normal"]),
                x["estado"],
                x["cupo_restante"],
                x["obtenidos"],
                x["canjeados"],
            ]
            for x in r["cupones"]
        ]
        partes.append(Table(filas, colWidths=[6.5 * cm, 2.5 * cm, 2 * cm, 2.5 * cm, 2.5 * cm], style=estilo_tabla))
    else:
        partes.append(Paragraph("Este negocio no tiene cupones.", est["Normal"]))
    doc.build(partes)
    return salida.getvalue()
